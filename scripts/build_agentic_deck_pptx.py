#!/usr/bin/env python3
"""Build the agentic-edge VP deck as a .pptx, mirroring docs/agentic-edge-deck.html
(final, fable-audited content). Dark engineering theme, data tables, provenance tags.
Run: python3 scripts/build_agentic_deck_pptx.py  ->  docs/agentic-edge-deck.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

INK=RGBColor(0xEE,0xF3,0xF7); MUT=RGBColor(0x9F,0xB2,0xC2); BG=RGBColor(0x0F,0x1A,0x24)
PANEL=RGBColor(0x12,0x1D,0x27); ACC=RGBColor(0xF5,0xA6,0x32); LINE=RGBColor(0x24,0x38,0x4A)
MEAS=RGBColor(0x39,0xC2,0x88); DERIV=RGBColor(0xE0,0xAD,0x3E); SRC=RGBColor(0x8E,0xA0,0xAD)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]
SW,SH=prs.slide_width,prs.slide_height

def slide(hero=False):
    s=prs.slides.add_slide(BLANK)
    r=s.shapes.add_shape(1,0,0,SW,SH); r.fill.solid()
    r.fill.fore_color.rgb=RGBColor(0x08,0x0F,0x16) if hero else BG
    r.line.fill.background(); r.shadow.inherit=False
    s.shapes._spTree.remove(r._element); s.shapes._spTree.insert(2,r._element)
    return s

def tb(s,x,y,w,h):
    b=s.shapes.add_textbox(x,y,w,h); b.text_frame.word_wrap=True; return b.text_frame

def para(tf,txt,size,color=INK,bold=False,first=False,space=6,align=PP_ALIGN.LEFT,font="Calibri"):
    p=tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment=align; p.space_after=Pt(space)
    # txt may be list of (run_text, color, bold) tuples, or a string
    runs = txt if isinstance(txt,list) else [(txt,color,bold)]
    for i,(t,c,b) in enumerate(runs):
        r=p.add_run(); r.text=t; r.font.size=Pt(size); r.font.bold=b
        r.font.color.rgb=c; r.font.name=font
    return p

def kicker(s,txt):
    tf=tb(s,Inches(0.7),Inches(0.5),Inches(12),Inches(0.4))
    para(tf,txt.upper(),13,ACC,bold=True,first=True)

def chip(label):
    c={"m":("MEASURED",MEAS),"d":("DERIVED",DERIV),"s":("SOURCED",SRC)}[label]
    return (f"  [{c[0]}]",c[1],True)

def snum(s,n):
    tf=tb(s,Inches(11.8),Inches(0.4),Inches(1.3),Inches(0.4))
    para(tf,f"{n} / 09",11,MUT,first=True,align=PP_ALIGN.RIGHT)

def table(s,x,y,w,rows,colw,header=True):
    nr,nc=len(rows),len(rows[0])
    gt=s.shapes.add_table(nr,nc,x,y,w,Inches(0.4*nr)).table
    for ci,cw in enumerate(colw): gt.columns[ci].width=cw
    for ri,row in enumerate(rows):
        for ci,val in enumerate(row):
            cell=gt.cell(ri,ci); cell.fill.solid(); cell.fill.fore_color.rgb=PANEL
            tf=cell.text_frame; tf.word_wrap=True; p=tf.paragraphs[0]
            rn=p.add_run(); rn.text=str(val); rn.font.size=Pt(12)
            rn.font.color.rgb=ACC if (ri>0 and ci==2) else INK
            rn.font.bold=(ri==0); rn.font.name="Consolas" if ci>=2 else "Calibri"
            if ri==0: rn.font.color.rgb=MUT; rn.font.size=Pt(10)
    return gt

# ---------- SLIDE 1 — BLUF ----------
s=slide(hero=True); snum(s,"01")
kicker(s,"BLUF · agentic workloads on edge silicon")
tf=tb(s,Inches(0.7),Inches(1.0),Inches(12),Inches(1.2))
para(tf,[("The agent's brain is ",INK,True),("bandwidth-bound",ACC,True),(", not TOPS-bound.",INK,True)],33,first=True)
tf=tb(s,Inches(0.7),Inches(2.55),Inches(12),Inches(0.9))
para(tf,"Measured on five devices from an RTX 5090 down to an NXP i.MX95 edge SoC: an on-device agent's cost is set by memory bandwidth, CPU, and model size — not by the TOPS number on the box.",18,MUT,first=True)
cards=[("01  Generation rides on bandwidth","Decode re-streams weights per token — memory-bound. Achieved BW holds at 55-70% of rated peak while TOPS varies many-fold. Armored headline: ~1.9x decode cost per model tier.",[chip("m"),chip("d")]),
("02  The NPU's job is the front-end","The NPU doesn't speed decode (bandwidth-bound on CPU/GPU/NPU alike). It earns its area on prefill/perception — on i.MX95 the Neutron offloads prefill 8.4x.",[chip("s")]),
("03  Model size ~ a 2x cost","+1 tier = ~1.9x decode + ~2x memory, invariant across three NVIDIA tiers. The biggest brains stop fitting the smallest boards.",[chip("m"),chip("d")]),
("04  So the sizing question flips","Not 'how many TOPS' but 'which tier can this board hold and feed' — then budget DDR+CPU for the loop, NPU TOPS for perception.",[chip("d")])]
for i,(h,body,chips) in enumerate(cards):
    cx=Inches(0.7+ (i%2)*6.1); cy=Inches(3.6+(i//2)*1.9)
    box=s.shapes.add_shape(1,cx,cy,Inches(5.9),Inches(1.8)); box.fill.solid(); box.fill.fore_color.rgb=PANEL
    box.line.color.rgb=LINE; box.shadow.inherit=False
    tf=box.text_frame; tf.word_wrap=True; tf.margin_left=Inches(0.2); tf.margin_top=Inches(0.12)
    para(tf,h,15,ACC,bold=True,first=True,space=4)
    para(tf,body,12,MUT,space=3)
    para(tf,[c for c in chips],10)

# ---------- SLIDE 2 — a/b framing ----------
s=slide(); snum(s,"02"); kicker(s,"definition · two very different 'agents'")
tf=tb(s,Inches(0.7),Inches(1.0),Inches(12),Inches(0.7)); para(tf,"Not all 'agentic' wants the same silicon.",24,INK,True,first=True)
tf=tb(s,Inches(0.7),Inches(1.7),Inches(12),Inches(0.5)); para(tf,"The word hides a 100x hardware split. Size for the one you're actually shipping.",16,MUT,first=True)
for i,(tag,h,body,on) in enumerate([
 ("(a) ships on-device today","Bounded skill-orchestrator","A fixed menu of tools — retrieve, transcribe, summarize, read/write files, call an API. A small on-device LLM (7B-class) is a capable enough brain. This is what Google Gemini Nano and Apple Intelligence ship. [SOURCED]",True),
 ("(b) wants a frontier brain","Autonomous code-gen / open-ended","Writes and runs novel code, plans open-endedly. Local-first in principle, but needs a large model to be good — today a cloud-scale brain, not an edge part.",False)]):
    box=s.shapes.add_shape(1,Inches(0.7+i*6.1),Inches(2.5),Inches(5.9),Inches(2.6)); box.fill.solid(); box.fill.fore_color.rgb=PANEL; box.line.color.rgb=ACC if on else LINE; box.shadow.inherit=False
    tf=box.text_frame; tf.word_wrap=True; tf.margin_left=Inches(0.25); tf.margin_top=Inches(0.2)
    para(tf,tag.upper(),12,MUT,bold=True,first=True,space=4); para(tf,h,18,INK,True,space=8); para(tf,body,13,MUT)
tf=tb(s,Inches(0.7),Inches(5.4),Inches(12),Inches(1.2))
para(tf,"This deck benchmarks (a) — the shippable case — using a local skill-orchestrator agent ('Skippy', Qwen 2.5 7B). 7 of 10 tasks are drawn from a general agent framework's capability set; 3 edge-perception tasks (transcribe, OCR, meeting->summary) were added.",13,MUT,first=True)

# ---------- SLIDE 3 — method ----------
s=slide(); snum(s,"03"); kicker(s,"method · 10 tasks, distinct resource profiles")
tf=tb(s,Inches(0.7),Inches(1.0),Inches(12),Inches(0.7)); para(tf,"Different agentic tasks hit different silicon.",24,INK,True,first=True)
rows=[["task","engine","dominant resource"],
["email · spec Q&A · doc brief","LLM + RAG","decode / bandwidth"],
["web search + summarize","web + LLM","network I/O + decode"],
["file read->summarize->write","file tools + LLM","filesystem + decode"],
["transcribe (ASR)","Whisper","audio compute"],
["OCR (image->text)","Tesseract","CPU — GPU idle (~3%)"],
["run sandboxed script","shell","process / compute"],
["multi-tool chain","agent loop","orchestration"]]
table(s,Inches(0.7),Inches(2.15),Inches(10),rows,[Inches(3.6),Inches(2.6),Inches(3.8)])
tf=tb(s,Inches(0.7),Inches(5.6),Inches(12),Inches(1.2))
para(tf,"Coherence gate: a task's resource numbers only count if its output is coherent — a broken output is a cheaper, unrepresentative computation ('broken is faster'). Every output was read and graded. [MEASURED]",13,MUT,first=True)

# ---------- SLIDE 4 — physics ----------
s=slide(); snum(s,"04"); kicker(s,"the core physics")
tf=tb(s,Inches(0.7),Inches(1.0),Inches(12),Inches(0.7)); para(tf,"Decode is memory-bandwidth-bound.",24,INK,True,first=True)
tf=tb(s,Inches(0.7),Inches(1.9),Inches(6),Inches(3.5))
para(tf,"Each generated token re-reads the whole model's weights from memory. Generation speed is set by how fast you stream weights, not peak math. Prefill (first token) IS compute-bound and parallel — a different resource.",14,INK,first=True,space=10)
para(tf,"• Generation-heavy tasks: ~85-90% of time in decode. [MEASURED]",13,MUT,space=5)
para(tf,"• RAG/tool tasks: 51-70% decode, +37-47% CPU/DDR orchestration. [MEASURED]",13,MUT,space=5)
para(tf,"• Achieved BW = 55-70% of each chip's rated peak — the memory-bound signature; tok/s does not track TOPS. [DERIVED]",13,MUT)
box=s.shapes.add_shape(1,Inches(7.1),Inches(1.9),Inches(5.5),Inches(3.0)); box.fill.solid(); box.fill.fore_color.rgb=PANEL; box.line.color.rgb=LINE; box.shadow.inherit=False
tf=box.text_frame; tf.word_wrap=True; tf.margin_left=Inches(0.25); tf.margin_top=Inches(0.2)
para(tf,"cross-instrument corroboration — Hexagon NPU",12,MUT,bold=True,first=True,space=6)
para(tf,'"Decode is DDR-bandwidth-bound. The scarce resource for LLM/VLA decode is DDR bandwidth, not TOPS."',15,INK,space=8)
para(tf,"sourced · qualcomm · SA8775P · 2026-09-15 [SOURCED]",11,SRC)

# ---------- SLIDE 5 — ladder (the key data slide) ----------
s=slide(); snum(s,"05"); kicker(s,"the hardware ladder · 7B-class Q4, decode tok/s")
tf=tb(s,Inches(0.7),Inches(1.0),Inches(12),Inches(0.7))
para(tf,"Decode tracks bandwidth, not TOPS.",24,INK,True,first=True)
rows=[["device","class","decode t/s","peak BW","prov."],
["RTX 5090","Blackwell dGPU","217.9","1792","MEAS"],
["Jetson AGX Thor","sm_110 Blackwell","40.9","273","MEAS"],
["Jetson AGX Orin","sm_87 Ampere","27.9","204.8","MEAS"],
["IQ-9075 (SA8775P)","Hexagon NPU · diff quant","9.45","—","SRC"],
["i.MX95 · A55 CPU","no accelerator","1.85","—","MEAS"],
["i.MX95 · +ARA240","+ M.2 accelerator","6.3","—","MEAS (dated)"]]
table(s,Inches(0.7),Inches(2.25),Inches(8.2),rows,[Inches(2.3),Inches(2.6),Inches(1.4),Inches(1.0),Inches(1.3)])
tf=tb(s,Inches(0.7),Inches(5.3),Inches(12),Inches(1.9))
para(tf,"The discriminator: Thor's rated AI throughput is many x Orin's, yet decode is only ~1.5x faster — tracking the 1.3x bandwidth gap (273 vs 205), not TOPS; achieved BW = 55-70% of rated peak [DERIVED].",13,MUT,first=True)
para(tf,"The ~118x top-to-bottom span is cross-instrument (5090 full-stack vs llama-bench vs iq9 different quant+runtime) — order-of-magnitude placement; the clean test is the within-board size ratio (next).",13,MUT)
para(tf,"i.MX95, honestly: bare A55 CPU 1.85 t/s; with the ARA240 accelerator 6.3 t/s, ~ NXP's published 6.51 spec — the accelerator makes a 7B interactive.",13,INK,True)

# ---------- SLIDE 6 — size tradeoff ----------
s=slide(); snum(s,"06"); kicker(s,"the sizing tradeoff · 7B vs 14B")
tf=tb(s,Inches(0.7),Inches(1.0),Inches(12),Inches(0.7))
para(tf,"A bigger brain costs ~2x — the same everywhere.",24,INK,True,first=True)
rows=[["board","7B dec","14B dec","ratio"],["RTX 5090","217.9","113.0","1.93x"],
["Thor","40.9","20.9","1.96x"],["Orin","27.9","14.6","1.91x"],["param ratio (ref)","","","1.94x"]]
table(s,Inches(0.7),Inches(2.15),Inches(6),rows,[Inches(2.4),Inches(1.2),Inches(1.2),Inches(1.2)])
tf=tb(s,Inches(0.7),Inches(4.6),Inches(6),Inches(2))
para(tf,"Decode penalty ~ parameter ratio, invariant across three NVIDIA tiers (Blackwell dGPU -> Ampere edge) — weight-streaming. Rates [MEASURED]; ratios [DERIVED].",13,MUT,first=True)
box=s.shapes.add_shape(1,Inches(7.1),Inches(1.9),Inches(5.5),Inches(3.2)); box.fill.solid(); box.fill.fore_color.rgb=PANEL; box.line.color.rgb=ACC; box.shadow.inherit=False
tf=box.text_frame; tf.word_wrap=True; tf.margin_left=Inches(0.25); tf.margin_top=Inches(0.2)
para(tf,"The real limit is fit, not speed",14,MUT,bold=True,first=True)
para(tf,"Capacity + context, not TOPS.",18,INK,True,space=8)
para(tf,"iq9: deployed NPU bundle caps a prompt at 256 tokens — longer RAG -> ~108s TTFT on CPU fallback [SOURCED]. A 14B on i.MX95 is untested + impractical (~1 t/s) [DERIVED].",12,MUT,space=6)
para(tf,"Sizing an edge agent is fit-then-feed — capacity + context + bandwidth, not the TOPS rating.",12,INK,True)

# ---------- SLIDE 7 — two-judge ----------
s=slide(); snum(s,"07"); kicker(s,"does the smaller brain actually work?")
tf=tb(s,Inches(0.7),Inches(1.0),Inches(12),Inches(0.7))
para(tf,"More model -> fewer fabrications.",24,INK,True,first=True)
tf=tb(s,Inches(0.7),Inches(1.75),Inches(12),Inches(0.9))
para(tf,"Two independent cross-family judges (Claude Sonnet + GPT-4o). Agreed: 4 PASS, 2 PARTIAL, 1 FAIL; 3 disagreements flagged, never averaged. Per-judge tallies bracket it — Sonnet 5/4/1, GPT-4o 4/3/3.  [MEASURED, two-judge]",14,MUT,first=True)
tf=tb(s,Inches(0.7),Inches(2.9),Inches(6),Inches(3.5))
para(tf,"4 PASS · 2 PARTIAL · 1 FAIL · 3 split",17,ACC,True,first=True,font="Consolas")
para(tf,"Clean passes: grounded/perception/tool tasks (file I/O, transcribe, OCR, run-script). Clean fail: multi-tool chain (narrates instead of calling tools). Partials: web-summarize (generic), meeting-summarize (fabricated attendees).",13,INK,space=8)
para(tf,"Both judges caught spec fabrication even when fluent — '256 MACs' (it's 1024), 'LPDDR4X' (it's LPDDR5X). Two splits PARTIAL-vs-FAIL; one PASS-vs-PARTIAL.",13,MUT)
box=s.shapes.add_shape(1,Inches(7.1),Inches(2.9),Inches(5.5),Inches(3.2)); box.fill.solid(); box.fill.fore_color.rgb=PANEL; box.line.color.rgb=ACC; box.shadow.inherit=False
tf=box.text_frame; tf.word_wrap=True; tf.margin_left=Inches(0.25); tf.margin_top=Inches(0.2)
para(tf,"the size fix · single-reader, two-judge pending",13,MUT,bold=True,first=True)
para(tf,"On 14B the same tasks (single-reader, not yet two-judge graded): it invokes the tool instead of narrating, and flags bad input instead of confabulating. [PENDING]",13,INK,space=8)
para(tf,"Neither is perfect — 14B still over-reads ('four Neutron NPUs'); size reduces, doesn't eliminate, fabrication. The point: an accuracy gate, not just a latency budget.",12,MUT)

# ---------- SLIDE 8 — NPU perception ----------
s=slide(); snum(s,"08"); kicker(s,"where the NPU pays off")
tf=tb(s,Inches(0.7),Inches(1.0),Inches(12),Inches(0.7)); para(tf,"The accelerator's job is perception, not generation.",24,INK,True,first=True)
tf=tb(s,Inches(0.7),Inches(1.9),Inches(6),Inches(3.8))
para(tf,"The NPU doesn't speed the generation loop — decode is bandwidth-bound on CPU, GPU and NPU alike. Where it earns its area is the compute-bound perception front-end:",14,INK,first=True,space=10)
para(tf,"• Prefill / TTFT — on i.MX95 the Neutron NPU offloads prefill 8.4x (@L=512) while decode stays a CPU wash. The measured NXP proof-point. [SOURCED]",13,MUT,space=5)
para(tf,"• Vision / VLM / VLA — image prefill dominates, decode tiny (mirror image of the LLM tasks).",13,MUT,space=5)
para(tf,"• ASR and other parallel, compute-bound perception.",13,MUT,space=8)
para(tf,"So an agentic edge SoC wants two separate budgets: DDR+CPU for the loop; NPU TOPS for perception + TTFT.",13,INK,True)
box=s.shapes.add_shape(1,Inches(7.1),Inches(1.9),Inches(5.5),Inches(3.4)); box.fill.solid(); box.fill.fore_color.rgb=PANEL; box.line.color.rgb=LINE; box.shadow.inherit=False
tf=box.text_frame; tf.word_wrap=True; tf.margin_left=Inches(0.25); tf.margin_top=Inches(0.2)
para(tf,"next (post-Friday): an on-edge VLA task",12,MUT,bold=True,first=True,space=6)
para(tf,'"A cat\'s in the yard — find it, show me a picture."',16,INK,True,space=8)
para(tf,"A perception->reason->act benchmark on a real 24/7 camera corpus (377 clips, free CV baseline). Early result: a 7B VLA structurally beats the motion baseline on STILL animals it physically can't see. Qwen2.5-VL-7B runs on the Hexagon NPU (qualcomm). [SOURCED]",12,MUT)

# ---------- SLIDE 9 — the ask ----------
s=slide(hero=True); snum(s,"09"); kicker(s,"the decision")
tf=tb(s,Inches(0.7),Inches(1.0),Inches(12),Inches(0.9))
para(tf,[("Size the ",INK,True),("subsystem",ACC,True),(", not the TOPS.",INK,True)],30,first=True)
tf=tb(s,Inches(0.7),Inches(2.0),Inches(12),Inches(0.8))
para(tf,"A design input for the next i.MX agentic part: size the subsystem (bandwidth, capacity, accelerator) to the model tier the product needs, and aim Neutron TOPS at perception.",16,MUT,first=True)
cards=[("MEASURED ON NXP SILICON","i.MX95 does a 7B — with its accelerator. Bare A55 CPU 1.85 t/s; +ARA240 6.3 t/s (~ NXP's 6.51 spec); Neutron offloads prefill 8.4x. Size the accelerator in, not out."),
("DESIGN FLOORS · 7B TIER","A decode path delivering >=~100 GB/s to the engine that runs it (CPU alone on i.MX95 reaches ~9), >=~10 GB DRAM for weights+KV, an accelerator for prefill. TOPS sized to perception."),
("PRODUCT RULE","A 7B orchestrator ships on small boards — passes grounded/perception/tool tasks cleanly; open-ended synthesis wants the next tier. 14B is a Thor-and-up part."),
("WATCH","The iq9's 256-token NPU prompt cap is the binding RAG constraint on the smallest parts — longer context silently falls to CPU (~108s TTFT). Size the context window deliberately.")]
for i,(h,body) in enumerate(cards):
    cx=Inches(0.7+(i%2)*6.1); cy=Inches(3.0+(i//2)*2.0)
    box=s.shapes.add_shape(1,cx,cy,Inches(5.9),Inches(1.85)); box.fill.solid(); box.fill.fore_color.rgb=RGBColor(0x14,0x22,0x2e); box.line.color.rgb=LINE; box.shadow.inherit=False
    tf=box.text_frame; tf.word_wrap=True; tf.margin_left=Inches(0.2); tf.margin_top=Inches(0.12)
    para(tf,h,12,ACC,bold=True,first=True,space=4)
    para(tf,body,12,MUT)

prs.save("docs/agentic-edge-deck.pptx")
print("wrote docs/agentic-edge-deck.pptx ·", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
